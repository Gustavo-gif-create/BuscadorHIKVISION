import streamlit as st
import pandas as pd
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="Buscador Hikvision Pro", layout="wide", page_icon="🔍")

# Nombre del archivo Excel oficial
EXCEL_FILE = 'Hikvision_Pricelist_DIAMOND_2026-10-01.xlsx'

# Carga y preparación de datos desde Excel
@st.cache_data
def load_data():
    # Lee directamente la primera hoja del archivo Excel
    df = pd.read_excel(EXCEL_FILE)
    
    # Limpieza de nombres de columnas (quita espacios extra al inicio o final)
    df.columns = df.columns.str.strip()
    
    # Asegurar que Stock sea numérico
    if 'Stock' in df.columns:
        df['Stock'] = pd.to_numeric(df['Stock'], errors='coerce').fillna(0)
    else:
        df['Stock'] = 0

    # Clasificación precisa por Macro Tipo de Producto
    def clasificar_tipo(row):
        cat = str(row.get('Cat. madre', '')).upper()
        line = str(row.get('Line Focus', '')).upper()
        mod = str(row.get('Modelo', '')).upper()
        
        if 'SWITCH' in line or 'NETWORKING' in cat or 'ROUTER' in line:
            return 'Switches / Networking'
        elif 'NVR' in line or 'NVR' in mod or 'NVR' in cat:
            return 'NVR'
        elif 'DVR' in line or 'DVR' in mod or ('BACK-END' in cat and ('DVR' in mod or 'XVR' in mod)):
            return 'DVR / XVR'
        elif 'PTZ' in line or 'PTZ' in mod:
            return 'Cámaras PTZ'
        elif 'TURBO HD' in line or mod.startswith('DS-2CE') or mod.startswith('THC') or 'TURBO' in mod:
            return 'Cámaras Análogas (Turbo HD)'
        elif 'P-FRONT-END' in cat or 'D-FRONT-END' in cat or 'IPC' in line or 'CAMERA' in line or 'CAM' in line or mod.startswith('DS-2CD') or mod.startswith('HCI'):
            return 'Cámaras IP'
        elif 'ACCESS CONTROL' in cat:
            return 'Control de Acceso'
        elif 'INTERCOM' in cat or 'INTERCOM' in line:
            return 'Videoporteros / Intercom'
        elif 'ALARM' in cat:
            return 'Alarmas'
        elif 'HIKMICRO' in cat or 'THERMAL' in line:
            return 'Térmicas / Hikmicro'
        elif 'ACCESSORY' in line or 'HOUSING' in line or 'BRAKET' in line or 'BRACKET' in line:
            return 'Accesorios y Montajes'
        elif 'DISPLAY' in cat or 'LED' in cat or 'MONITOR' in line:
            return 'Monitores / Pantallas / LED'
        else:
            return 'Otros / Varios'

    df['Tipo de Producto'] = df.apply(clasificar_tipo, axis=1)
    return df

try:
    df = load_data()
except Exception as e:
    st.error(f"No se pudo cargar el archivo '{EXCEL_FILE}'. Asegúrate de haberlo subido al repositorio en GitHub.")
    st.stop()

st.title("🔍 Buscador y Selector Comercial Hikvision")
st.caption(f"Base de datos activa: **{len(df):,}** artículos")

# --- BARRA LATERAL: FILTROS DINÁMICOS Y CONSECUENTES ---
st.sidebar.header("🎯 Filtros Generales")

df_curr = df.copy()

# 1. Tipo de Producto
tipos_producto = ['Todos'] + sorted(df_curr['Tipo de Producto'].unique().tolist())
tipo_sel = st.sidebar.selectbox("Tipo de Producto", tipos_producto)

if tipo_sel != 'Todos':
    df_curr = df_curr[df_curr['Tipo de Producto'] == tipo_sel]

# 2. Marca
marcas = ['Todas'] + sorted(df_curr['Marca'].dropna().unique().tolist()) if 'Marca' in df_curr.columns else ['Todas']
marca_sel = st.sidebar.selectbox("Marca", marcas)

if marca_sel != 'Todas':
    df_curr = df_curr[df_curr['Marca'] == marca_sel]

# 3. Categoría Madre
cats = ['Todas'] + sorted(df_curr['Cat. madre'].dropna().unique().tolist()) if 'Cat. madre' in df_curr.columns else ['Todas']
cat_sel = st.sidebar.selectbox("Categoría Madre", cats)

if cat_sel != 'Todas':
    df_curr = df_curr[df_curr['Cat. madre'] == cat_sel]

# 4. Line Focus
lines = ['Todas'] + sorted(df_curr['Line Focus'].dropna().unique().tolist()) if 'Line Focus' in df_curr.columns else ['Todas']
line_sel = st.sidebar.selectbox("Línea (Line Focus)", lines)

if line_sel != 'Todas':
    df_curr = df_curr[df_curr['Line Focus'] == line_sel]

# 5. Segmento
segmentos = ['Todos'] + sorted(df_curr['Segmento'].dropna().unique().tolist()) if 'Segmento' in df_curr.columns else ['Todos']
seg_sel = st.sidebar.selectbox("Segmento", segmentos)

if seg_sel != 'Todos':
    df_curr = df_curr[df_curr['Segmento'] == seg_sel]

# 6. Filtro de Stock
solo_stock = st.sidebar.checkbox("Mostrar solo con Stock disponible (>0)")
if solo_stock and 'Stock' in df_curr.columns:
    df_curr = df_curr[df_curr['Stock'] > 0]

# --- FILTROS POR CARACTERÍSTICAS TÉCNICAS DINÁMICAS ---
st.sidebar.markdown("---")
st.sidebar.header("🛠️ Características Técnicas")

diccionario_tags = {
    "Resolución / Calidad": ["1080P", "2MP", "3K", "4MP", "5MP", "8MP", "4K", "HD"],
    "Funciones / Red / IA": ["PoE", "Gigabit", "SFP", "AcuSense", "ColorVu", "DarkFighter", "WDR", "Microphone", "Audio", "Built-in MIC", "Deep learning", "Human/Vehicle classification", "Face Recognition", "GPS", "Wi-Fi", "4G", "BSD", "ADAS"],
    "Protección / Construcción": ["IP67", "IK10", "Water resistant", "Vandal", "Metal"],
    "Compresión / Visión": ["H.265+", "H.264", "EXIR", "Hybrid light", "Smart Light", "Dual Light"]
}

desc_col = df_curr['Descripción'].astype(str) if 'Descripción' in df_curr.columns else pd.Series(['']*len(df_curr))
mod_col = df_curr['Modelo'].astype(str) if 'Modelo' in df_curr.columns else pd.Series(['']*len(df_curr))

texto_acumulado = (desc_col.str.cat(sep=" ") + " " + mod_col.str.cat(sep=" ")).lower()

for cat_tech, tags in diccionario_tags.items():
    tags_validos = [tag for tag in tags if tag.lower() in texto_acumulado]
    if tags_validos:
        seleccionados = st.sidebar.multiselect(f"{cat_tech}:", tags_validos, key=f"multi_{cat_tech}")
        for kw in seleccionados:
            mask_d = df_curr['Descripción'].astype(str).str.contains(kw, case=False, na=False) if 'Descripción' in df_curr.columns else False
            mask_m = df_curr['Modelo'].astype(str).str.contains(kw, case=False, na=False) if 'Modelo' in df_curr.columns else False
            df_curr = df_curr[mask_d | mask_m]

# Búsqueda libre
busqueda_libre = st.sidebar.text_input("O escribe una palabra clave específica:")
if busqueda_libre:
    mask_m = df_curr['Modelo'].astype(str).str.contains(busqueda_libre, case=False, na=False) if 'Modelo' in df_curr.columns else False
    mask_s = df_curr['SAP'].astype(str).str.contains(busqueda_libre, case=False, na=False) if 'SAP' in df_curr.columns else False
    mask_d = df_curr['Descripción'].astype(str).str.contains(busqueda_libre, case=False, na=False) if 'Descripción' in df_curr.columns else False
    df_curr = df_curr[mask_m | mask_s | mask_d]

# --- VISTA PRINCIPAL ---
st.metric("Productos encontrados", len(df_curr))

precios_cols = ['Diamond', 'Black', 'Ruby', 'Platinum', 'Gold Gremio', 'Gold Plus', 'Gold Prime', 'MAP ARS', 'MAP Web']
cols_a_mostrar = ['SAP', 'Modelo', 'Marca', 'Tipo de Producto', 'Cat. madre', 'Line Focus', 'Segmento', 'Status', 'Stock'] + precios_cols + ['Descripción']
cols_existentes = [c for c in cols_a_mostrar if c in df_curr.columns]

st.dataframe(
    df_curr[cols_existentes],
    use_container_width=True,
    height=500
)

# --- GENERAR EXCEL PRESENTACIÓN ---
def generar_excel_estetico(dataframe):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lista de Precios Hikvision"
    ws.views.sheetView[0].showGridLines = True

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    border_thin = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    for r in dataframe_to_rows(dataframe[cols_existentes], index=False, header=True):
        ws.append(r)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.font = data_font
            cell.border = border_thin
            cell.alignment = Alignment(vertical="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()

excel_data = generar_excel_estetico(df_curr)

st.download_button(
    label="📊 Descargar Excel Presentación (.xlsx)",
    data=excel_data,
    file_name="Lista_Precios_Hikvision.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
